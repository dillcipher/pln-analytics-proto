"""Bounded-memory CUSTOMER_LOCATION ingestion for fixed TO masters."""
from __future__ import annotations

import logging
import os
import shutil
import sqlite3
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.application.etl.etl_orchestrator import ETLOrchestrator
from app.etl.merger.monthly_merger import MonthlyMerger
from app.etl.validator.validator import DatasetValidator

logger = logging.getLogger(__name__)
_INSTALLED = False


def _chunks(path: Path, dataset: str, n: int):
    sheet = DatasetValidator.get_sheet_name(path, dataset)
    header = DatasetValidator.detect_header_row(path, sheet)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        values = next(
            ws.iter_rows(
                min_row=header + 1,
                max_row=header + 1,
                values_only=True,
            ),
            None,
        )
        if not values:
            return
        columns = [
            str(v).strip().upper() if v is not None else ""
            for v in values
        ]
        batch = []
        for row in ws.iter_rows(
            min_row=header + 2,
            values_only=True,
        ):
            batch.append(tuple(row[:len(columns)]))
            if len(batch) >= n:
                yield pd.DataFrame.from_records(
                    batch,
                    columns=columns,
                )
                batch.clear()
        if batch:
            yield pd.DataFrame.from_records(
                batch,
                columns=columns,
            )
    finally:
        wb.close()


def _text(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def _stream_masters(
    files: list[Path],
    output: Path,
    chunk_rows: int,
) -> Path:
    import pyarrow as pa
    import pyarrow.parquet as pq

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    staging = (
        output.parent.parent
        / ".customer_location_stream"
        / uuid.uuid4().hex
    )
    staging.mkdir(parents=True, exist_ok=True)
    temp = staging / output.name

    db = sqlite3.connect(staging / "selected.sqlite3")
    db.execute(
        """
        CREATE TABLE selected (
            seq INTEGER PRIMARY KEY AUTOINCREMENT,
            idpel TEXT UNIQUE NOT NULL,
            unitupi TEXT,
            unitap TEXT,
            unitup TEXT,
            x REAL,
            y REAL,
            valid INTEGER NOT NULL,
            priority INTEGER NOT NULL
        )
        """
    )
    db.commit()

    transformer = MonthlyMerger.TRANSFORMERS["CUSTOMER_LOCATION"]

    try:
        for source in files:
            source = Path(source)
            priority = MonthlyMerger._coordinate_source_priority(source.name)
            logger.info(
                "STREAMING CUSTOMER_LOCATION MASTER | file=%s | priority=%s | chunk_rows=%s",
                source.name,
                priority,
                chunk_rows,
            )

            for number, frame in enumerate(
                _chunks(source, "CUSTOMER_LOCATION", chunk_rows),
                start=1,
            ):
                if frame.empty:
                    continue
                frame.columns = frame.columns.map(str).str.strip().str.upper()
                frame = transformer.transform(frame)
                if frame is None or frame.empty or "IDPEL" not in frame.columns:
                    continue

                for column in (
                    "UNITUPI",
                    "UNITAP",
                    "UNITUP",
                    "KOORDINAT_X",
                    "KOORDINAT_Y",
                ):
                    if column not in frame.columns:
                        frame[column] = ""

                frame["IDPEL"] = MonthlyMerger._normalize_idpel_series(frame["IDPEL"])
                frame = frame.loc[frame["IDPEL"].ne("")].copy()
                if frame.empty:
                    continue

                x = pd.to_numeric(frame["KOORDINAT_X"], errors="coerce")
                y = pd.to_numeric(frame["KOORDINAT_Y"], errors="coerce")
                valid = (x.notna() & y.notna()).astype(int)

                rows = [
                    (
                        _text(i),
                        _text(a),
                        _text(b),
                        _text(c),
                        float(xx) if pd.notna(xx) else None,
                        float(yy) if pd.notna(yy) else None,
                        int(v),
                        int(priority),
                    )
                    for i, a, b, c, xx, yy, v in zip(
                        frame["IDPEL"].tolist(),
                        frame["UNITUPI"].tolist(),
                        frame["UNITAP"].tolist(),
                        frame["UNITUP"].tolist(),
                        x.tolist(),
                        y.tolist(),
                        valid.tolist(),
                    )
                ]

                db.executemany(
                    """
                    INSERT INTO selected(
                        idpel, unitupi, unitap, unitup,
                        x, y, valid, priority
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(idpel) DO UPDATE SET
                        unitupi=excluded.unitupi,
                        unitap=excluded.unitap,
                        unitup=excluded.unitup,
                        x=excluded.x,
                        y=excluded.y,
                        valid=excluded.valid,
                        priority=excluded.priority
                    WHERE
                        excluded.valid > selected.valid
                        OR (
                            excluded.valid = selected.valid
                            AND excluded.priority > selected.priority
                        )
                    """,
                    rows,
                )
                db.commit()
                logger.info(
                    "STREAMING CUSTOMER_LOCATION PART | file=%s | chunk=%s | rows=%s",
                    source.name,
                    number,
                    len(rows),
                )
                del frame, x, y, valid, rows

        writer = None
        total = 0
        for frame in pd.read_sql_query(
            """
            SELECT
                idpel AS IDPEL,
                unitupi AS UNITUPI,
                unitap AS UNITAP,
                unitup AS UNITUP,
                x AS KOORDINAT_X,
                y AS KOORDINAT_Y
            FROM selected
            ORDER BY seq
            """,
            db,
            chunksize=10_000,
        ):
            frame["DATASET"] = "CUSTOMER_LOCATION"
            table = pa.Table.from_pandas(frame, preserve_index=False)
            if writer is None:
                writer = pq.ParquetWriter(temp, table.schema, compression="snappy")
            elif table.schema != writer.schema:
                table = table.cast(writer.schema, safe=False)
            writer.write_table(table)
            total += len(frame)

        if writer is None:
            raise ValueError("CUSTOMER_LOCATION master streaming produced no rows.")
        writer.close()
        writer = None
        os.replace(temp, output)
        logger.info(
            "STREAMING CUSTOMER_LOCATION MASTER COMPLETE | output=%s | rows=%s",
            output,
            total,
        )
        return output
    finally:
        try:
            if "writer" in locals() and writer is not None:
                writer.close()
        except Exception:
            logger.exception("CUSTOMER_LOCATION writer cleanup failed")
        db.close()
        shutil.rmtree(staging, ignore_errors=True)


def install_streaming_customer_location_patch() -> None:
    global _INSTALLED
    if _INSTALLED:
        return

    original_merge = MonthlyMerger.merge
    cache: dict[tuple[str, ...], Path] = {}

    def patched_merge(dataset, month, files, output_dir):
        if dataset != "CUSTOMER_LOCATION" or month is None:
            return original_merge(dataset, month, files, output_dir)

        normalized = [
            ETLOrchestrator._normalize_filename(Path(path).name)
            for path in files
        ]
        if not files or not all(
            value in ETLOrchestrator.COORDINATE_MASTER_FILES
            for value in normalized
        ):
            return original_merge(dataset, month, files, output_dir)

        key = tuple(sorted(str(Path(path).resolve()) for path in files))
        target = (
            Path(output_dir)
            / "customer_location"
            / f"customer_location_{month}.parquet"
        )
        base = cache.get(key)
        if base is None or not base.exists():
            base = _stream_masters(
                files=files,
                output=target,
                chunk_rows=max(
                    100,
                    int(os.getenv("CUSTOMER_LOCATION_STREAM_CHUNK_ROWS", "250")),
                ),
            )
            cache[key] = base
            return base

        if base.resolve() == target.resolve():
            return target

        target.parent.mkdir(parents=True, exist_ok=True)
        target.unlink(missing_ok=True)
        try:
            os.link(base, target)
            logger.info(
                "REUSE CUSTOMER_LOCATION MASTER VIA HARDLINK | month=%s | source=%s",
                month,
                base.name,
            )
        except OSError:
            shutil.copy2(base, target)
            logger.info(
                "REUSE CUSTOMER_LOCATION MASTER VIA COPY | month=%s | source=%s",
                month,
                base.name,
            )
        return target

    MonthlyMerger.merge = staticmethod(patched_merge)
    _INSTALLED = True
    logger.info("Installed bounded-memory CUSTOMER_LOCATION master patch.")
