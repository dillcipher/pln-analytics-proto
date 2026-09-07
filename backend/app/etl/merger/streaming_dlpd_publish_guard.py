"""Memory-safety guards for large ETL jobs."""
from __future__ import annotations

import gc
import hashlib
import logging
import os
import re
import shutil
import sqlite3
import uuid
from pathlib import Path

import pandas as pd

import app.etl.merger.streaming_dlpd_merger_patch as streaming
from app.application.etl.etl_orchestrator import ETLOrchestrator
from app.etl.merger.monthly_merger import MonthlyMerger
from app.etl.transformers.anev_transformer import ANEVTransformer
from app.etl.validator.validator import DatasetValidator

logger = logging.getLogger(__name__)
_INSTALLED = False
DLPD = {"DLPD_PASCABAYAR", "DLPD_PRABAYAR"}


def _month_from_path(path: Path) -> str:
    p = path.stem.split("_")
    return p[-2] if len(p) >= 3 and p[-1].startswith("part") else ""


def _chunks(path: Path, dataset: str, n: int):
    """
    Stream an Excel sheet in a single forward pass.

    The previous implementation repeatedly called pandas.ExcelFile.parse() with
    an ever-growing skiprows range. On multi-million-row ANEV files this becomes
    progressively more expensive and can appear to freeze after millions of
    rows because each next chunk re-scans the workbook from the beginning.
    openpyxl read_only mode advances through the worksheet once instead.
    """
    from openpyxl import load_workbook

    sheet = DatasetValidator.get_sheet_name(path, dataset)
    header = DatasetValidator.detect_header_row(path, sheet)

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        worksheet = workbook[sheet]
        rows = worksheet.iter_rows(values_only=True)

        # Skip everything before the detected header, then read it once.
        for _ in range(header):
            next(rows, None)

        header_row = next(rows, None)
        if header_row is None:
            return

        cols = [
            str(value).strip() if value is not None else ""
            for value in header_row
        ]

        batch = []
        for row in rows:
            # Trim trailing empty cells while preserving interior columns.
            values = list(row[: len(cols)])
            if len(values) < len(cols):
                values.extend([None] * (len(cols) - len(values)))
            batch.append(values)

            if len(batch) >= n:
                yield pd.DataFrame.from_records(batch, columns=cols)
                batch = []

        if batch:
            yield pd.DataFrame.from_records(batch, columns=cols)
    finally:
        workbook.close()
        del workbook
        gc.collect()

def _canonical_anev_columns(columns) -> list[str]:
    """Normalize Excel headers without destroying canonical names."""
    normalized: list[str] = []
    seen: dict[str, int] = {}

    for index, value in enumerate(columns):
        name = str(value).replace("\n", " ").replace("\r", " ").strip().upper()
        name = re.sub(r"\s+", "_", name)
        name = re.sub(r"[^A-Z0-9_]", "_", name)
        name = re.sub(r"_+", "_", name).strip("_")
        if not name:
            name = f"UNNAMED_{index}"

        count = seen.get(name, 0)
        seen[name] = count + 1
        if count:
            name = f"{name}_{count}"

        normalized.append(name)

    return normalized


def _parse_anev_dates(series: pd.Series) -> pd.Series:
    """Fast bounded parser; avoids dateutil fallback on every chunk."""
    if pd.api.types.is_datetime64_any_dtype(series):
        return series

    numeric = pd.to_numeric(series, errors="coerce")
    numeric_mask = numeric.notna() & numeric.between(1, 100000)
    result = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")

    if numeric_mask.any():
        result.loc[numeric_mask] = pd.to_datetime(
            numeric.loc[numeric_mask],
            unit="D",
            origin="1899-12-30",
            errors="coerce",
        )

    text_mask = ~numeric_mask & series.notna()
    if text_mask.any():
        text = series.loc[text_mask].astype(str).str.strip()
        result.loc[text_mask] = pd.to_datetime(
            text,
            errors="coerce",
            format="mixed",
        )

    return result


def _anev_key_series(chunk: pd.DataFrame) -> pd.Series:
    """
    Build a bounded, stable deduplication key.

    ANEV is not required to contain IDPEL. Older code unconditionally called
    drop_duplicates(["IDPEL"]), which crashes valid ANEV files that only carry
    LOCATION_CODE/READ_DATE/SUSPECT_NAME.
    """
    preferred = [
        column
        for column in ("IDPEL", "LOCATION_CODE", "READ_DATE", "SUSPECT_NAME")
        if column in chunk.columns
    ]
    columns = preferred or list(chunk.columns)

    if not columns:
        return pd.Series(dtype="object", index=chunk.index)

    values = chunk.loc[:, columns].copy()
    for column in columns:
        values[column] = values[column].fillna("").astype(str).str.strip()

    joined = values.astype(str).agg("\x1f".join, axis=1)
    return joined.map(
        lambda value: hashlib.sha1(value.encode("utf-8", "ignore")).hexdigest()
    )


def _filter_fresh_anev_rows(
    db: sqlite3.Connection,
    chunk: pd.DataFrame,
) -> pd.DataFrame:
    """Keep only first-seen rows without assuming any particular ANEV key column."""
    keys = _anev_key_series(chunk)
    if keys.empty:
        return chunk.iloc[0:0].copy()

    # First occurrence in the current bounded chunk wins.
    first_mask = ~keys.duplicated(keep="first")
    unique_keys = keys.loc[first_mask].tolist()
    if not unique_keys:
        return chunk.iloc[0:0].copy()

    existing: set[str] = set()
    # SQLite commonly limits bound parameters to 999. Keep batches comfortably below.
    for start in range(0, len(unique_keys), 900):
        batch = unique_keys[start : start + 900]
        placeholders = ",".join("?" for _ in batch)
        existing.update(
            row[0]
            for row in db.execute(
                f"SELECT dedup_key FROM processed_keys "
                f"WHERE dedup_key IN ({placeholders})",
                batch,
            )
        )

    fresh_keys = [key for key in unique_keys if key not in existing]
    if not fresh_keys:
        return chunk.iloc[0:0].copy()

    db.executemany(
        "INSERT OR IGNORE INTO processed_keys(dedup_key) VALUES (?)",
        ((key,) for key in fresh_keys),
    )

    # Do not force a full SQLite fsync/commit for every 1,000-row chunk.
    # The caller owns a short transaction window and commits periodically.
    fresh_set = set(fresh_keys)
    keep_mask = first_mask & keys.isin(fresh_set)
    return chunk.loc[keep_mask].copy()


def _align_anev_schema(
    chunk: pd.DataFrame,
    schema_columns: list[str] | None,
) -> tuple[pd.DataFrame, list[str]]:
    """
    Keep one deterministic Parquet schema across streamed files.

    Missing columns are added as nulls; unexpected later columns are retained
    only when the first schema has not yet been established.

    Historically `schema_columns` started as None and got fixed to whichever
    chunk was processed FIRST for a given month -- a real bug (see
    `_prescan_anev_schema`'s docstring): if that first chunk came from a
    source file with an incomplete/different column set (e.g. a header
    mis-detected on a summary row, or simply a different source workbook
    template than other files feeding the same month), every later chunk --
    including ones with the real, complete column set -- got silently
    truncated down to that first, possibly-wrong schema via
    `chunk.reindex()` here. `_stream_anev` now always pre-scans every
    source file's header before writing anything and passes a fully
    resolved `schema_columns` in from the start, so this function's only
    remaining job is the per-chunk reindex/null-fill against that
    already-correct, already-complete target -- there is no longer a
    "first chunk wins" moment.
    """
    if schema_columns is None:
        return chunk, list(chunk.columns)

    aligned = chunk.reindex(columns=schema_columns)
    return aligned, schema_columns


def _prescan_anev_schema(files: list[Path]) -> list[str]:
    """Resolve one complete, deterministic column set for a month's ANEV
    output BEFORE streaming any data, by reading just each source file's
    header row (cheap -- header-only, not the millions of data rows below
    it).

    Root cause this replaces: `_stream_anev` used to leave
    `schema_columns` as `None` until the very first data chunk arrived,
    then locked onto that chunk's columns for the rest of the month --
    across every remaining chunk AND every remaining file. When a month is
    assembled from more than one source workbook (`files: list[Path]`,
    plural) and those workbooks don't share an identical column layout
    (different preparer, different template, a header mis-detected on a
    stray summary/pivot row above the real table -- confirmed live
    2026-09-01: this is what produced anev_<month>.parquet files whose
    only columns were things like UNITAP/UNITUP/UNITUPI/CURRENT_N,
    unrelated to the real per-row ANEV fields), whichever file's data
    happened to stream first silently decided the schema for the whole
    month, and every column that ONLY appeared in a later file was
    dropped for that entire month -- not just for that one file's rows.
    Pre-scanning every file's header first and taking the union (in
    first-seen order) means the month's output always has every real
    column the source data actually contains, regardless of which file
    streams first.

    MONTH and DATASET are not header columns (they're assigned
    programmatically once per chunk, see `_stream_anev`) but must be part
    of the target schema from the start too, or the same "whichever chunk
    established the schema first" problem would just recur one level up
    for them specifically.
    """
    from openpyxl import load_workbook

    columns: list[str] = []
    seen: set[str] = set()

    def _add(name: str) -> None:
        if name not in seen:
            seen.add(name)
            columns.append(name)

    _add("MONTH")
    _add("DATASET")

    for source in files:
        try:
            sheet = DatasetValidator.get_sheet_name(source, "ANEV")
            header = DatasetValidator.detect_header_row(source, sheet)
            workbook = load_workbook(
                filename=source,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            try:
                worksheet = workbook[sheet]
                rows = worksheet.iter_rows(values_only=True)
                for _ in range(header):
                    next(rows, None)
                header_row = next(rows, None)
                if header_row is None:
                    continue
                raw = [
                    str(value).strip() if value is not None else ""
                    for value in header_row
                ]
                for name in _canonical_anev_columns(raw):
                    _add(name)
            finally:
                workbook.close()
                del workbook
        except Exception:
            logger.exception(
                "ANEV schema pre-scan failed for %s; its columns will "
                "only be included if another file in this month shares "
                "them.",
                source,
            )

    gc.collect()
    return columns


def _stream_anev(month: str, files: list[Path], output_dir: Path) -> Path:
    import pyarrow as pa
    import pyarrow.parquet as pq

    out = Path(output_dir) / "anev" / f"anev_{month}.parquet"
    out.parent.mkdir(parents=True, exist_ok=True)

    tmpdir = Path(output_dir) / ".anev_stream_staging" / uuid.uuid4().hex
    tmpdir.mkdir(parents=True, exist_ok=True)
    tmp = tmpdir / out.name

    db = sqlite3.connect(tmpdir / "seen.sqlite3")
    db.execute("PRAGMA journal_mode=OFF")
    db.execute("PRAGMA synchronous=OFF")
    db.execute("PRAGMA temp_store=MEMORY")
    db.execute(
        "CREATE TABLE IF NOT EXISTS processed_keys("
        "dedup_key TEXT PRIMARY KEY)"
    )

    writer = None
    writer_schema = None
    # Resolved BEFORE any data streams (see _prescan_anev_schema's
    # docstring) so the "whichever chunk/file processes first quietly
    # decides the month's schema forever" bug can't happen. header-only
    # reads are cheap; this does not load any of the actual data rows.
    schema_columns: list[str] | None = _prescan_anev_schema(files)
    total = 0
    skipped_without_rows = 0
    chunk_count = 0
    rows_since_commit = 0
    n = max(250, int(os.getenv("ANEV_STREAM_CHUNK_ROWS", "500")))
    commit_every_rows = max(
        n,
        int(os.getenv("ANEV_DEDUP_COMMIT_ROWS", "50000")),
    )
    transformer = ANEVTransformer()

    # Confirmed live 2026-09-01: writing one Parquet row group per ~500-1000
    # row ETL chunk (the previous behavior -- one writer.write_table() call
    # per chunk) produces hundreds of tiny row groups per month for a
    # dataset this size. Measured locally: reconciling that many row
    # groups' footer statistics across 6 monthly files at DuckDB
    # CREATE VIEW time costs meaningfully more resident memory than the
    # same total data written as far fewer, larger row groups (~25%
    # higher peak RSS in a like-for-like local comparison: 300 row
    # groups/file vs. 12 row groups/file over the same 300K rows) -- a
    # plausible contributor to the OOM crashes observed building
    # fact_anev on the 512MB host, on top of union_by_name's own
    # multi-file reconciliation cost. The small ANEV_STREAM_CHUNK_ROWS
    # value is there to bound in-*memory* pandas/pyarrow processing per
    # chunk, which is a separate concern from how large a row group ends
    # up on *disk* -- so several processed chunks are buffered here and
    # flushed together as one larger row group, without changing how
    # many rows are ever held as a pandas DataFrame at once.
    row_group_target_rows = max(
        n,
        int(os.getenv("ANEV_PARQUET_ROW_GROUP_ROWS", "20000")),
    )
    pending_tables: list = []
    pending_rows = 0
    # Tracks the schema new tables get cast to as they're appended to the
    # CURRENT pending batch (reset after every flush) -- separate from
    # writer_schema (which is fixed for the whole file once the very first
    # row group is written). Needed because pa.concat_tables() requires
    # every input to share an identical schema, and two chunks with the
    # same column NAMES can still infer different pyarrow TYPES (e.g. a
    # column that's all-null in one chunk vs. real floats in another --
    # confirmed by a synthetic test mixing a narrow, mostly-empty source
    # file with a wide one in the same batch).
    pending_schema = None

    def _flush_pending() -> None:
        nonlocal writer, writer_schema, pending_tables, pending_rows, pending_schema, total
        if not pending_tables:
            return
        import pyarrow as pa

        table = (
            pending_tables[0]
            if len(pending_tables) == 1
            else pa.concat_tables(pending_tables)
        )
        pending_schema = None
        if writer is None:
            writer_schema = table.schema
            writer = pq.ParquetWriter(
                tmp,
                writer_schema,
                compression="snappy",
            )
        elif table.schema != writer_schema:
            table = table.cast(writer_schema, safe=False)

        writer.write_table(table)
        total += pending_rows
        logger.info(
            "STREAMING ANEV ROW GROUP FLUSH | month=%s | rows=%s | total=%s",
            month,
            pending_rows,
            total,
        )
        pending_tables = []
        pending_rows = 0

    try:
        for source in files:
            logger.info(
                "STREAMING ANEV SOURCE | %s | chunk_rows=%s",
                source.name,
                n,
            )

            for chunk in _chunks(Path(source), "ANEV", n):
                try:
                    chunk.columns = _canonical_anev_columns(chunk.columns)

                    # IDPEL cleaning is optional for ANEV. It is applied when present,
                    # but valid ANEV chunks without IDPEL must continue normally.
                    if "IDPEL" in chunk.columns:
                        chunk = transformer.clean_idpel(chunk)

                    if chunk.empty:
                        skipped_without_rows += 1
                        continue

                    logger.debug(
                        "ANEV DEDUP START | month=%s | input_rows=%s",
                        month,
                        len(chunk),
                    )
                    chunk = _filter_fresh_anev_rows(db, chunk)
                    chunk_count += 1
                    rows_since_commit += len(chunk)
                    if rows_since_commit >= commit_every_rows:
                        db.commit()
                        rows_since_commit = 0
                        logger.info(
                            "ANEV DEDUP CHECKPOINT | month=%s | chunks=%s | rows=%s",
                            month,
                            chunk_count,
                            total,
                        )
                    if chunk.empty:
                        continue

                    if "READ_DATE" in chunk.columns:
                        chunk["READ_DATE"] = _parse_anev_dates(
                            chunk["READ_DATE"]
                        )

                    for column in transformer.NUMERIC_COLUMNS:
                        if column in chunk.columns:
                            chunk[column] = pd.to_numeric(
                                chunk[column],
                                errors="coerce",
                            )

                    if "DATASET" not in chunk.columns:
                        chunk["DATASET"] = "ANEV"

                    # Confirmed live 2026-09-01: this streaming ANEV path
                    # (installed in app/main.py, replacing the legacy
                    # MonthlyMerger.merge() path which always did
                    # `merged["MONTH"] = month`) never wrote a MONTH
                    # column at all -- every anev_<month>.parquet file
                    # this path produces is missing it entirely. The
                    # Executive Dashboard's `WHERE MONTH IS NOT NULL`
                    # query failed with `Binder Error: Referenced column
                    # "MONTH" not found` as a direct result. Source data
                    # has no MONTH column to preserve (unlike DLPD, whose
                    # source rows can carry their own per-row month), so
                    # this always assigns the month this file is being
                    # streamed for -- mirrors the DATASET assignment right
                    # above.
                    if "MONTH" not in chunk.columns:
                        chunk["MONTH"] = month

                    for column in chunk.columns:
                        if pd.api.types.is_object_dtype(chunk[column]):
                            chunk[column] = (
                                chunk[column]
                                .fillna("")
                                .astype(str)
                                .str.strip()
                            )

                    chunk, schema_columns = _align_anev_schema(
                        chunk,
                        schema_columns,
                    )
                    logger.debug(
                        "ANEV PARQUET CONVERT START | month=%s | rows=%s",
                        month,
                        len(chunk),
                    )
                    table = pa.Table.from_pandas(
                        chunk,
                        preserve_index=False,
                    )

                    logger.debug(
                        "ANEV PARQUET BUFFER | month=%s | rows=%s | pending=%s",
                        month,
                        len(chunk),
                        pending_rows + len(chunk),
                    )
                    if pending_schema is None:
                        pending_schema = table.schema
                    elif table.schema != pending_schema:
                        table = table.cast(pending_schema, safe=False)
                    pending_tables.append(table)
                    pending_rows += len(chunk)
                    if pending_rows >= row_group_target_rows:
                        _flush_pending()

                    logger.info(
                        "STREAMING ANEV PART | month=%s | rows=%s | total=%s",
                        month,
                        len(chunk),
                        total + pending_rows,
                    )
                finally:
                    del chunk

            gc.collect()

        # Flush the final bounded dedup transaction before publishing output.
        db.commit()

        # Any buffered rows smaller than one full row-group target still
        # need to reach disk -- most likely the tail end of the month's
        # last file.
        _flush_pending()

        if writer is None:
            raise ValueError(
                f"ANEV streaming produced no rows for {month}; "
                f"empty_chunks={skipped_without_rows}"
            )

        writer.close()
        writer = None
        os.replace(tmp, out)

        logger.info(
            "STREAMING ANEV COMPLETE | month=%s | rows=%s",
            month,
            total,
        )
        return out
    finally:
        if writer is not None:
            writer.close()
        db.close()
        shutil.rmtree(tmpdir, ignore_errors=True)
        gc.collect()


def _patch_orchestrator() -> None:
    original_resolve = ETLOrchestrator._resolve_dlpd_month_cache.__func__
    original_expand = ETLOrchestrator._expand_processing_groups.__func__

    def resolve_hints(cls, grouped, job_folder):
        hints = {str(m) for m in cls._get_business_months(grouped) if m}
        cache = {}
        for (dataset, _), files in grouped.items():
            if dataset not in DLPD:
                continue
            for r in files:
                if cls._is_valid_file(r) and r.get("filename"):
                    cache.setdefault(r["filename"], set(hints))
        if cache:
            logger.info(
                "DLPD MONTH RESOLUTION DEFERRED | files=%s | hints=%s",
                len(cache),
                sorted(hints),
            )
            return cache
        return original_resolve(cls, grouped, job_folder)

    ETLOrchestrator._resolve_dlpd_month_cache = classmethod(resolve_hints)
    ETLOrchestrator._expand_processing_groups = classmethod(
        lambda cls, grouped, job_folder, dlpd_month_cache=None:
        original_expand(cls, grouped, job_folder, dlpd_month_cache)
    )

    original_merge = MonthlyMerger.merge
    cache = {}

    def safe_merge(dataset, month, files, output_dir):
        if dataset == "ANEV" and month is not None:
            return _stream_anev(str(month), files, output_dir)

        if dataset != "CUSTOMER_LOCATION" or month is None:
            return original_merge(dataset, month, files, output_dir)

        masters = [
            p for p in files
            if ETLOrchestrator._normalize_filename(p.name)
            in ETLOrchestrator.COORDINATE_MASTER_FILES
        ]

        if len(masters) != len(files) or not masters:
            return original_merge(dataset, month, files, output_dir)

        key = tuple(
            sorted(str(Path(p).resolve()) for p in masters)
        )
        target = (
            Path(output_dir)
            / "customer_location"
            / f"customer_location_{month}.parquet"
        )
        target.parent.mkdir(parents=True, exist_ok=True)

        base = cache.get(key)
        if base is None or not base.exists():
            base = original_merge(dataset, month, masters, output_dir)
            cache[key] = base

        if base.resolve() != target.resolve():
            shutil.copy2(base, target)

        logger.info(
            "REUSE FIXED COORDINATE MASTER | month=%s | source=%s",
            month,
            base.name,
        )
        return target

    MonthlyMerger.merge = staticmethod(safe_merge)


def install_streaming_dlpd_publish_guard() -> None:
    global _INSTALLED
    if _INSTALLED:
        return

    def safe_publish(
        output_dir: Path,
        dataset: str,
        staged_files: list[Path],
    ):
        folder = Path(output_dir) / "dlpd"
        folder.mkdir(parents=True, exist_ok=True)

        prefix = (
            "dlpd_pascabayar_"
            if dataset == "DLPD_PASCABAYAR"
            else "dlpd_prabayar_"
        )
        months = {
            m for m in (_month_from_path(p) for p in staged_files) if m
        }

        if not months:
            raise RuntimeError(
                f"Cannot publish DLPD {dataset}: no monthly partitions"
            )

        prepared = []
        for p in staged_files:
            final = folder / p.name
            hidden = folder / f".{p.name}.new"
            os.replace(p, hidden)
            prepared.append((hidden, final))

        for m in months:
            for old in folder.glob(f"{prefix}{m}_part*.parquet"):
                old.unlink(missing_ok=True)
            # Also remove any legacy, pre-partition file for this month
            # (e.g. dlpd_prabayar_202606.parquet, no "_part" suffix) --
            # written by the old app.etl.large_dlpd_stream merge path.
            # The glob above never matches that name, so before this fix
            # a stale legacy file could sit in storage forever even after
            # a later run genuinely republished that month: the
            # warehouse's broad "dlpd_<dataset>*.parquet" glob picks up
            # both, and if the legacy file is corrupt/incomplete (e.g. a
            # MONTH-only stub) it can stand in for -- or mix into -- the
            # month it claims to be. Confirmed live 2026-09-02:
            # dlpd_prabayar_202606.parquet (legacy-named, MONTH-only) was
            # the sole file ever hydrated for DLPD Prabayar.
            legacy = folder / f"{prefix}{m}.parquet"
            if legacy.exists():
                legacy.unlink(missing_ok=True)

        published = {}
        try:
            for hidden, final in prepared:
                os.replace(hidden, final)
                m = _month_from_path(final)
                if m:
                    published.setdefault(m, final)
        except Exception:
            for hidden, _ in prepared:
                hidden.unlink(missing_ok=True)
            raise

        logger.info(
            "DLPD PUBLISH COMPLETE | dataset=%s | months=%s | files=%s",
            dataset,
            sorted(months),
            len(prepared),
        )
        return published

    streaming._publish_staged_outputs = safe_publish
    _patch_orchestrator()
    _INSTALLED = True
    logger.info(
        "Installed memory-safe DLPD + ANEV orchestration guard."
    )
