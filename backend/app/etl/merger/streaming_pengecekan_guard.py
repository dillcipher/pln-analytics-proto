"""Memory-bounded PENGECEKAN ingestion for small production containers."""
from __future__ import annotations

import logging
import os
import shutil
import sqlite3
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.etl.merger.monthly_merger import MonthlyMerger
from app.etl.transformers.pengecekan_transformer import PengecekanTransformer
from app.etl.validator.validator import DatasetValidator

logger = logging.getLogger(__name__)
_INSTALLED = False


def _iter_excel_chunks(filepath: Path, dataset: str, chunk_rows: int):
    """Read XLSX rows incrementally without materialising the workbook."""
    filepath = Path(filepath)
    sheet = DatasetValidator.get_sheet_name(filepath, dataset)
    header = DatasetValidator.detect_header_row(filepath, sheet)
    workbook = load_workbook(filepath, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet]
        header_values = next(
            worksheet.iter_rows(
                min_row=header + 1,
                max_row=header + 1,
                values_only=True,
            ),
            None,
        )
        if not header_values:
            return

        columns = [
            str(value).strip().upper() if value is not None else ""
            for value in header_values
        ]
        batch: list[tuple] = []
        for row in worksheet.iter_rows(
            min_row=header + 2,
            values_only=True,
        ):
            batch.append(tuple(row[: len(columns)]))
            if len(batch) >= chunk_rows:
                yield pd.DataFrame.from_records(batch, columns=columns)
                batch.clear()
        if batch:
            yield pd.DataFrame.from_records(batch, columns=columns)
    finally:
        workbook.close()


def _align_chunk_schema(table: "pa.Table", target_schema: "pa.Schema") -> "pa.Table":
    """Cast one PENGECEKAN chunk's table onto the writer's locked schema.

    Chunks are read independently, so pandas can infer a different dtype for
    the same column across chunks -- e.g. chunk 1 has real numbers in a
    column (-> float64/double), but a later chunk has a stray blank/non-
    numeric cell for that same column, which pandas/pyarrow instead infers
    as string. A plain `table.cast(target_schema, safe=False)` then raises
    `pyarrow.lib.ArrowInvalid: Failed to parse string: '' as a scalar of
    type double` (confirmed live 2026-09-01, first GitHub Actions ETL run --
    and confirmed AGAIN on a second run after a first attempt at this fix
    only nulled out exact `''` values: this environment runs pandas 3.0,
    whose new default string dtype is not object dtype, so
    `pd.api.types.is_object_dtype()` upstream in _stream_pengecekan no
    longer reliably flags every text column for the fillna("")/str
    normalization pass that used to guarantee 'missing' meant exactly the
    string ''; some other unparseable value can reach this cast instead).
    Try the direct cast first -- cheap, and correct for the overwhelming
    majority of chunks where nothing actually changed dtype. Only on a
    genuine parse failure, fall back to a lenient pandas-based coercion
    that turns whatever specific value can't be parsed into a real null,
    for numeric AND temporal target types alike, instead of crashing the
    whole PENGECEKAN merge over one bad cell in one chunk.
    """
    import pyarrow as pa
    import pandas as pd

    arrays = []
    for field in target_schema:
        if field.name in table.column_names:
            column = table.column(field.name)
        else:
            column = pa.nulls(table.num_rows, type=field.type)

        if column.type != field.type:
            try:
                column = column.cast(field.type, safe=False)
            except pa.lib.ArrowInvalid:
                series = column.to_pandas()
                if pa.types.is_floating(field.type) or pa.types.is_integer(field.type):
                    series = pd.to_numeric(series, errors="coerce")
                    if pa.types.is_integer(field.type):
                        # Confirmed live 2026-09-02: a later chunk can carry
                        # a genuine fractional value (e.g. "0.910") for a
                        # column an earlier chunk's dtype inference locked
                        # as int64 in writer.schema. pa.array(..., type=
                        # <int64>) below raises ArrowInvalid ("Float value
                        # ... was truncated converting to int64") on that --
                        # an unhandled exception here previously escaped
                        # _align_chunk_schema entirely, quarantining the
                        # WHOLE PENGECEKAN dataset over one bad cell in one
                        # chunk, exactly what this function's lenient-
                        # coercion fallback exists to prevent. Null out only
                        # the values that don't round-trip losslessly into
                        # the locked integer type, instead of a lossy
                        # truncate/crash.
                        non_integral = series.notna() & (series % 1 != 0)
                        if non_integral.any():
                            logger.warning(
                                "PENGECEKAN CHUNK SCHEMA COERCION DROPPED "
                                "NON-INTEGRAL VALUES | field=%s locked as %s "
                                "| count=%s | sample=%s",
                                field.name,
                                field.type,
                                int(non_integral.sum()),
                                series[non_integral].head(5).tolist(),
                            )
                            series = series.where(~non_integral, other=pd.NA)
                elif pa.types.is_temporal(field.type):
                    series = pd.to_datetime(series, errors="coerce")
                else:
                    series = series.where(series.notna(), None).astype(object)
                column = pa.array(series, type=field.type, from_pandas=True)

        arrays.append(column)

    return pa.Table.from_arrays(arrays, schema=target_schema)


def _stream_pengecekan(files: list[Path], output_dir: Path) -> Path:
    """Stream all PENGECEKAN sources into one bounded-memory parquet."""
    if not files:
        raise ValueError("No PENGECEKAN files supplied.")

    output_dir = Path(output_dir)
    final_dir = output_dir / "pengecekan"
    final_dir.mkdir(parents=True, exist_ok=True)
    final_path = final_dir / "pengecekan.parquet"

    staging_root = output_dir / ".pengecekan_stream_staging"
    staging_dir = staging_root / uuid.uuid4().hex
    staging_dir.mkdir(parents=True, exist_ok=True)
    temp_path = staging_dir / "pengecekan.parquet"
    seen_db = sqlite3.connect(staging_dir / "seen_idpel.sqlite3")
    writer = None
    total = 0
    chunk_rows = max(250, int(os.getenv("PENGECEKAN_STREAM_CHUNK_ROWS", "500")))

    try:
        seen_db.execute("CREATE TABLE seen_idpel (idpel TEXT PRIMARY KEY)")
        seen_db.commit()

        import pyarrow as pa
        import pyarrow.parquet as pq

        for source in files:
            logger.info(
                "STREAMING PENGECEKAN SOURCE | file=%s | size=%s | chunk_rows=%s",
                Path(source).name,
                Path(source).stat().st_size,
                chunk_rows,
            )
            for chunk_number, chunk in enumerate(
                _iter_excel_chunks(Path(source), "PENGECEKAN", chunk_rows),
                start=1,
            ):
                if chunk.empty:
                    continue

                chunk.columns = chunk.columns.map(str).str.strip().str.upper()
                chunk["SOURCE_FILE"] = Path(source).name
                transformed = PengecekanTransformer().transform(chunk)
                if transformed.empty:
                    continue

                if "IDPEL" in transformed.columns:
                    ids = transformed["IDPEL"].fillna("").astype(str).str.strip()
                    transformed = transformed.loc[ids.ne("")].copy()
                    if transformed.empty:
                        continue
                    ids = transformed["IDPEL"].astype(str).str.strip()
                    unique_ids = ids.drop_duplicates().tolist()
                    fresh_ids: list[str] = []
                    for value in unique_ids:
                        if seen_db.execute(
                            "SELECT 1 FROM seen_idpel WHERE idpel = ?",
                            (value,),
                        ).fetchone() is None:
                            seen_db.execute(
                                "INSERT INTO seen_idpel(idpel) VALUES (?)",
                                (value,),
                            )
                            fresh_ids.append(value)
                    seen_db.commit()
                    if not fresh_ids:
                        continue
                    transformed = transformed.loc[ids.isin(set(fresh_ids))].copy()

                for column in transformed.columns:
                    # pandas 3.0's default text dtype is its own "str"
                    # dtype, not object dtype -- is_object_dtype() alone
                    # misses it, so a genuinely-text column with a blank
                    # cell in this chunk could keep its NA representation
                    # unnormalized instead of becoming "" like every other
                    # chunk's text columns. Catch both: is_object_dtype for
                    # a genuinely mixed-type column (pandas can't type that
                    # as pure "str"), is_string_dtype for the new default.
                    is_textlike = pd.api.types.is_object_dtype(
                        transformed[column]
                    ) or pd.api.types.is_string_dtype(transformed[column])
                    if is_textlike:
                        transformed[column] = (
                            transformed[column].fillna("").astype(str).str.strip()
                        )

                table = pa.Table.from_pandas(transformed, preserve_index=False)
                if writer is None:
                    writer = pq.ParquetWriter(
                        temp_path,
                        table.schema,
                        compression="snappy",
                    )
                elif table.schema != writer.schema:
                    table = _align_chunk_schema(table, writer.schema)

                writer.write_table(table)
                total += len(transformed)
                logger.info(
                    "STREAMING PENGECEKAN PART | chunk=%s | rows=%s | total=%s",
                    chunk_number,
                    len(transformed),
                    total,
                )

        if writer is None:
            raise ValueError("Streaming PENGECEKAN produced no rows.")

        writer.close()
        writer = None
        os.replace(temp_path, final_path)
        logger.info(
            "STREAMING PENGECEKAN COMPLETE | rows=%s | output=%s",
            total,
            final_path,
        )
        return final_path
    finally:
        if writer is not None:
            writer.close()
        seen_db.close()
        shutil.rmtree(staging_dir, ignore_errors=True)
        try:
            staging_root.rmdir()
        except OSError:
            pass


def install_streaming_pengecekan_guard() -> None:
    global _INSTALLED
    if _INSTALLED:
        return

    original_merge = MonthlyMerger.merge

    def safe_merge(
        dataset: str,
        month: str | None,
        files: list[Path],
        output_dir: Path,
    ) -> Path:
        if dataset == "PENGECEKAN":
            if month is not None:
                raise ValueError("PENGECEKAN must be processed without a business month.")
            return _stream_pengecekan(files, output_dir)
        return original_merge(dataset, month, files, output_dir)

    MonthlyMerger.merge = staticmethod(safe_merge)
    _INSTALLED = True
    logger.info(
        "Installed memory-bounded PENGECEKAN merger guard | chunk_rows=%s",
        os.getenv("PENGECEKAN_STREAM_CHUNK_ROWS", "500"),
    )
